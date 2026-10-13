import sys

try:
    import openpyxl
    wb = openpyxl.load_workbook(r'C:\Users\Administrator\Downloads\BCT\Template\PTK_Bao cao tuan_MinhTQ.xlsx', data_only=True)
    print("Sheets:", wb.sheetnames)
    for name in wb.sheetnames:
        sheet = wb[name]
        print(f"\n=== Sheet: {name} ({sheet.max_row} rows, {sheet.max_column} cols) ===")
        for r in range(1, min(30, sheet.max_row + 1)):
            row_vals = [str(sheet.cell(r, c).value or '').strip() for c in range(1, min(20, sheet.max_column + 1))]
            if any(row_vals):
                print(f"Row {r:2d}: " + " | ".join(row_vals[:12]))
except Exception as e:
    print("Error:", e)
